"""Rebuild the WristMET training set from the public Zenodo "Phase 2 - Data".

Everything here follows the public upstream code
(https://github.com/HAbitsLab/WristBased-EE-Estimation, commit
3d3517cb95ef490d5e867ebf6a7207b8f7278788) and cites the file and lines it reproduces;
where the upstream code is not runnable as published or the public data differ from
what the code reads, the decision and its justification are noted next to the code
and collected in the build report.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sqlite3
import zipfile
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from ..features import (
    freq_intensity,
    rmssd_l1,
)
from ..preprocess import mirror_to_right_wrist, resample_linear

LOCAL_TZ = "America/Chicago"  # helper_preprocess.py:watch_add_datetime
SAMPLE_RATE_HZ = 20

#: in-lab.py:34 (the '1001': 'NA' entry is the participant excluded from the public data).
MET_CART_START = {
    "1000": "01/27/2022 13:43:15", "1002": "02/09/2022 10:30:45", "1003": "01/12/2022 15:47:36",
    "1004": "01/21/2022 13:49:59", "1005": "02/12/2022 11:57:42", "1006": "03/22/2022 10:23:05",
    "1007": "03/28/2022 12:49:21", "1008": "04/12/2022 12:11:31", "1009": "04/15/2022 12:37:46",
    "1010": "05/06/2022 11:14:05", "1011": "05/11/2022 12:02:56", "1012": "05/24/2022 16:42:47",
    "1013": "05/27/2022 11:34:26", "1014": "06/03/2022 11:50:35", "1015": "06/06/2022 10:41:15",
    "1016": "08/01/2022 12:47:00", "1017": "08/16/2022 13:11:40", "1018": "08/26/2022 11:54:56",
    "1019": "09/21/2022 11:41:39", "1020": "10/04/2022 11:17:21", "1021": "10/14/2022 13:52:46",
    "1022": "10/19/2022 18:31:47", "1023": "10/21/2022 11:42:49", "1024": "11/03/2022 11:40:41",
    "1025": "11/07/2022 10:46:36", "1026": "11/16/2022 11:07:17",
}

#: in-lab.py:266 / 480. Label 0 = sedentary, 1 = everything else (in-lab.py:314-317).
SEDENTARY_ACTIVITIES = (
    "Rest",
    "Typing on a computer while seated",
    "Reading a book or magazine while reclining",
    "Lying down while doing nothing",
)
#: in-lab.py:481-482 (regression training set = windows of these activities, in-lab.py:484-488).
LIGHT_ACTIVITIES = ("Walking 2 mph on treadmill", "Standing while fidgeting", "Sweeping slowly ")
VIGOROUS_ACTIVITIES = (
    "Walking 3.5 mph on treadmill",
    "Push-ups against the wall",
    "Squats (shoulder length legs, get down to 90 degree angle)",
    "General aerobics video",
    "Running 4 mph on a treadmill",
    "Chester Step Test (0.25 m step at a rate of 30 steps per minute)",
)

#: Activity -> sedentary / light / moderate-vigorous groups used for reporting only.
#: The paper's per-class RMSE bins are by measured METs (<1.5, 1.5-<3, >=3); see ``met_class``.
MET_CLASS_EDGES = (1.5, 3.0)


def met_class(met: np.ndarray) -> np.ndarray:
    """0 = sedentary (<1.5 METs), 1 = light (1.5-<3), 2 = moderate/vigorous (>=3)."""
    met = np.asarray(met, dtype=np.float64)
    return np.where(met < MET_CLASS_EDGES[0], 0, np.where(met < MET_CLASS_EDGES[1], 1, 2))


# --------------------------------------------------------------------------------------
# Demographics
# --------------------------------------------------------------------------------------

#: The demographic table that the upstream code builds from private lab spreadsheets
#: (in-lab.py:40-82) is printed in the public notebook ``Result Visualization.ipynb`` (cell 6 output,
#: MIT, same commit): participant -> (weight_kg, height_cm, gender, age). Weight = MetCart record (lb) x
#: 0.45359237, height = MetCart "H/W" record in inches (0.1 in resolution) x 2.54, age = year the notebook
#: was run (2024) minus birth year. P1001 is the participant excluded from the public data.
UPSTREAM_NOTEBOOK_DEMOGRAPHICS = {
    "1000": (126.098679, 166.878, 0, 41), "1001": (129.273825, 177.8, 1, 49),
    "1002": (104.326245, 163.83, 0, 27), "1003": (96.161582, 166.878, 0, 62),
    "1004": (128.820233, 176.53, 1, 44), "1005": (114.75887, 173.99, 1, 45),
    "1006": (105.23343, 175.768, 0, 33), "1007": (101.604691, 162.56, 0, 47),
    "1008": (86.636143, 160.528, 1, 31), "1009": (106.140615, 182.88, 1, 34),
    "1010": (84.821773, 160.02, 0, 35), "1011": (117.026831, 167.132, 0, 56),
    "1012": (75.749926, 157.48, 0, 51), "1013": (104.326245, 180.34, 1, 35),
    "1014": (105.687022, 165.608, 0, 38), "1015": (100.697506, 170.18, 0, 39),
    "1016": (91.625659, 152.4, 0, 47), "1017": (84.368181, 160.02, 0, 65),
    "1018": (124.737902, 170.18, 0, 38), "1019": (98.429544, 158.75, 0, 55),
    "1020": (108.408576, 176.53, 1, 50), "1021": (118.387609, 160.02, 0, 28),
    "1022": (89.811289, 160.02, 0, 63), "1023": (115.666054, 190.5, 1, 65),
    "1024": (85.275366, 158.75, 0, 53), "1025": (135.170526, 179.07, 1, 31),
    "1026": (101.151099, 176.53, 1, 46),
}

DEMOGRAPHICS_SOURCES = ("primary", "agd", "notebook")


@dataclass
class Demographics:
    """One participant. ``*_agd`` fields come from the ActiGraph headers in the Zenodo data.

    ``height_cm`` / ``age_years`` / ``weight_kg`` are the values used for the *primary* training
    features: sex, age at the visit and weight from the ActiGraph header, height from the upstream
    notebook table (MetCart height, 0.1 in resolution, the same values that reproduce the paper's
    Table 1: female mean height 64.3 in; the header heights are rounded to whole inches).
    """

    participant: str  # "1000"
    sex: str
    gender: int  # male = 1 (in-lab.py:64-72)
    age_years: int  # age at the visit (ActiGraph header) - reproduces the paper's Table 1 (female 44.2, male 41.3)
    height_cm: float
    weight_kg: float
    height_cm_agd: float
    weight_kg_gt3x: float | None
    height_cm_gt3x: float | None
    age_notebook: int  # upstream notebook: run year (2024) - birth year
    height_cm_notebook: float
    weight_kg_notebook: float
    hip_side: str
    hip_dominance: str
    left_handed: bool

    @property
    def bmi(self) -> float:
        return self.weight_kg / (self.height_cm / 100.0) ** 2

    def with_source(self, source: str) -> "Demographics":
        """Variant for sensitivity analyses: ``"agd"`` (all header values) or ``"notebook"`` (all notebook values)."""
        if source == "primary":
            return self
        if source == "agd":
            return replace(self, height_cm=self.height_cm_agd)
        if source == "notebook":
            return replace(
                self, height_cm=self.height_cm_notebook, weight_kg=self.weight_kg_notebook, age_years=self.age_notebook
            )
        raise ValueError(f"unknown demographics source {source!r}")


def read_demographics(pdir: Path) -> Demographics:
    """Sex / age / height / mass / hip side from the ActiGraph ``.agd`` ``settings`` table.

    The public data contain no ``Participant Measurement Record.xlsx`` / ``Participant Log.xlsx`` /
    ``Participant Genders.xlsx`` that ``in-lab.py:40-72`` reads, but each participant's ActiGraph
    files carry the same fields (the AGD mass is a whole number of pounds, as in the lab record).
    The ``.gt3x`` ``info.txt`` is read as a cross-check.
    """
    pid = pdir.name[1:]
    agd = next((pdir / "Actigraph").glob("*.agd"))
    con = sqlite3.connect(f"file:{agd}?mode=ro", uri=True)
    try:
        s = dict(con.execute("select settingName, settingValue from settings"))
    finally:
        con.close()
    sex = str(s["sex"]).strip()
    gt3x_mass = gt3x_height = None
    gt3x = next((pdir / "Actigraph").glob("*.gt3x"), None)
    if gt3x is not None:
        with zipfile.ZipFile(gt3x) as z:
            info = z.read("info.txt").decode("utf8", "replace")
        d = dict(re.findall(r"^([^:]+):\s*(.*?)\s*$", info, re.M))
        gt3x_mass = _to_float(d.get("Mass"))
        gt3x_height = _to_float(d.get("Height"))
    side = str(s["side"]).strip()
    dominance = str(s["dominance"]).strip()
    # The hip device is worn contralateral to the dominant hand (paper, Methods "ActiGraph
    # accelerometer"); every public participant has dominance "Non-Dominant" on the hip.
    if dominance.lower().startswith("non"):
        dominant_hand = "left" if side.lower().startswith("r") else "right"
    else:
        dominant_hand = side.lower()
    nb_weight, nb_height, nb_gender, nb_age = UPSTREAM_NOTEBOOK_DEMOGRAPHICS[pid]
    gender = 1 if sex[:1].upper() == "M" else 0
    if gender != nb_gender or abs(float(s["mass"]) - nb_weight) > 1e-5:
        raise ValueError(f"P{pid}: header sex/mass disagree with the upstream notebook table")
    return Demographics(
        participant=pid,
        sex=sex,
        gender=gender,
        age_years=int(float(s["age"])),
        height_cm=nb_height,
        weight_kg=float(s["mass"]),
        height_cm_agd=float(s["height"]),
        weight_kg_gt3x=gt3x_mass,
        height_cm_gt3x=gt3x_height,
        age_notebook=nb_age,
        height_cm_notebook=nb_height,
        weight_kg_notebook=nb_weight,
        hip_side=side,
        hip_dominance=dominance,
        left_handed=dominant_hand == "left",
    )


def _to_float(x: str | None) -> float | None:
    try:
        return float(x) if x not in (None, "", "None") else None
    except ValueError:
        return None


def handedness_notes(pdir: Path) -> list[str]:
    """Free-text cells of the activity log that mention handedness (evidence for the hip-side rule)."""
    wb = load_workbook(pdir / "Calorie Harmony Phase 2 Activity Log.xlsx", data_only=True)
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and re.search(r"left.?hand|right.?hand|lefty|handed|dominan", cell.value, re.I):
                    hits.append(f"{ws.title}!{cell.coordinate}: {cell.value.strip()[:160]}")
    return hits


def demographics_table(demos: list[Demographics], age_offset: float = 0.0, source: str = "primary") -> pd.DataFrame:
    """Upstream ``df_demographic`` (in-lab.py:74-82): gender/age/BMI and the interaction terms."""
    rows = []
    for d in (x.with_source(source) for x in demos):
        age = d.age_years + age_offset
        bmi = d.bmi
        rows.append(
            {
                "Participant": int(d.participant),
                "gender": d.gender,
                "age": age,
                "BMI": bmi,
                "weight": d.weight_kg,
                "height": d.height_cm,
                "gender_age": age * d.gender,
                "gender_BMI": bmi * d.gender,
                "age_BMI": age * bmi,
                "age_gender_BMI": age * d.gender * bmi,
                "left_handed": d.left_handed,
            }
        )
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------------------
# Activity log, MetCart, wrist data
# --------------------------------------------------------------------------------------

@dataclass
class LogRow:
    row: int  # 0-based index inside the 24 activity rows (df_logs index in upstream)
    activity: str
    include: object
    start: dt.time | None
    stop: dt.time | None


def read_activity_log(path: Path) -> list[LogRow]:
    """Rows 3-26 of the "Activity Log" sheet, columns looked up by header name.

    Upstream reads ``usecols="A:L", skiprows=1, nrows=24`` (in-lab.py:243); the column
    layout varies between participants (``Include`` is in column K, L or M), so the
    columns are located by their header text instead of by position.
    """
    ws = load_workbook(path, data_only=True)["Activity Log"]
    header = {str(ws.cell(2, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value is not None}
    out = []
    for i, r in enumerate(range(3, 27)):
        start = ws.cell(r, header["Start Time"]).value
        stop = ws.cell(r, header["Expected Stop Time"]).value
        out.append(
            LogRow(
                row=i,
                activity=ws.cell(r, 1).value,
                include=ws.cell(r, header["Include"]).value,
                start=start if isinstance(start, dt.time) else None,
                stop=stop if isinstance(stop, dt.time) else None,
            )
        )
    return out


def local_to_epoch_ms(naive: pd.Timestamp) -> int:
    return int(naive.tz_localize(LOCAL_TZ).value // 1_000_000)


@dataclass
class MetCart:
    t_ms: np.ndarray  # UTC epoch ms of each breath row
    met: np.ndarray
    n_lines_total: int
    n_lines_parsed: int
    n_data_like_lines: int  # lines that start with a clock time (HH:MM:SS)


_DATA_LINE = re.compile(r"^\d\d:\d\d:\d\d")


def read_metcart(path: Path, date_mmddyyyy: str) -> MetCart:
    """``helper_preprocess.get_metcart`` (lines 131-144), unchanged parsing rules.

    A line is a data row iff ``len(line.strip())`` is 72 or 80; the clock time is
    ``line.strip()[:8]`` and the METs are the fixed-width slice ``line.strip()[21:24]``
    (so a value of 10.0 METs or more is read as e.g. ``0.3`` - 13 such rows exist in the
    public data - and is kept, exactly as the upstream parser does). Rows whose slice is
    not numeric are dropped. The date comes from the hard-coded dictionary.
    """
    lines = path.read_text(errors="replace").splitlines()
    stamps, vals = [], []
    for line in lines:
        s = line.strip()
        if len(s) == 72 or len(s) == 80:
            stamps.append(date_mmddyyyy + " " + s[:8])
            vals.append(s[21:24])
    n_data_like = sum(1 for line in lines if _DATA_LINE.match(line.strip()))
    df = pd.DataFrame({"Datetime": stamps, "met": vals})
    df["Datetime"] = pd.to_datetime(df["Datetime"], format="%m/%d/%Y %H:%M:%S")
    df["met"] = pd.to_numeric(df["met"], errors="coerce")
    df = df[df["met"].notnull()]
    idx = pd.DatetimeIndex(df["Datetime"]).tz_localize(LOCAL_TZ).tz_convert("UTC").tz_localize(None)
    t_ms = ((idx - pd.Timestamp("1970-01-01")) // pd.Timedelta(milliseconds=1)).to_numpy(np.int64)
    return MetCart(t_ms, df["met"].to_numpy(np.float64), len(lines), len(stamps), n_data_like)


@dataclass
class WristStream:
    t_ms: np.ndarray  # int64 UTC epoch ms of the 20 Hz grid
    x: np.ndarray  # [n, 3]
    n_chunks: int
    n_selected_rows: int
    original_t_ms: np.ndarray  # selected original sample times (sorted), for gap statistics


def utc_date_of_local_start(session_start: str) -> dt.date:
    naive = pd.Timestamp(dt.datetime.strptime(session_start, "%m/%d/%Y %H:%M:%S"))
    return naive.tz_localize(LOCAL_TZ).tz_convert("UTC").date()


def read_wrist_stream(pdir: Path, sensor: str, session_start: str) -> WristStream:
    """Original chunks -> 20 Hz series, as ``wrist_preprocessing.py`` (lines 52-92) did for the shipped files.

    Chunks are concatenated, sorted by the numeric ``Time`` (ms), restricted to the UTC date of the
    MetCart start and the following calendar day, then resampled with the upstream ``resample``
    (vectorised in ``preprocess.resample_linear``; verified bit-for-bit against the shipped
    ``Resampled`` CSVs by ``verify_against_shipped``). Gaps are interpolated across (the shipped
    data used an infinite gap tolerance).
    """
    folder = pdir / "Wrist Data" / "Clean" / "Original" / sensor
    files = sorted(folder.glob("*.csv"))
    df = pd.concat([pd.read_csv(f) for f in files], ignore_index=True)
    cols = ["accX", "accY", "accZ"] if sensor == "Accelerometer" else [c for c in df.columns if c.startswith("rot")]
    df["Time"] = df["Time"].astype(np.float64)
    day1 = utc_date_of_local_start(session_start)
    day2 = day1 + dt.timedelta(days=1)
    dates = pd.to_datetime(df["Time"], unit="ms", utc=True).dt.date
    df = df[(dates == day1) | (dates == day2)]
    df = df.sort_values("Time", kind="stable").reset_index(drop=True)
    t, x = resample_linear(df["Time"].to_numpy(), df[cols].to_numpy(np.float64), SAMPLE_RATE_HZ, None)
    return WristStream(t.astype(np.int64), x, len(files), len(df), df["Time"].to_numpy(np.int64))


def verify_against_shipped(pdir: Path, sensor: str, stream: WristStream) -> dict:
    """Compare our re-resampled stream with the shipped ``Resampled`` CSV (prefix overlap)."""
    name = "Accelerometer/acc_resample.csv" if sensor == "Accelerometer" else "Gyroscope/gyro_resample.csv"
    shipped = pd.read_csv(pdir / "Wrist Data" / "Clean" / "Resampled" / name)
    n = min(len(shipped), stream.t_ms.shape[0])
    time_equal = bool(np.array_equal(shipped["Time"].to_numpy(np.int64)[:n], stream.t_ms[:n]))
    max_diff = float(np.max(np.abs(shipped.iloc[:n, 1:4].to_numpy(np.float64) - stream.x[:n])))
    return {
        "len_shipped": int(len(shipped)),
        "len_ours": int(stream.t_ms.shape[0]),
        "overlap": int(n),
        "time_equal": time_equal,
        "max_abs_diff": max_diff,
    }


# --------------------------------------------------------------------------------------
# Windowing
# --------------------------------------------------------------------------------------

@dataclass
class WindowDataset:
    window_s: int
    n_samples: int  # samples per full window
    meta: pd.DataFrame
    windows: np.ndarray  # [n, n_samples, 6]; (gyr_xyz, acc_xyz); NaN-padded where a window is short (unmirrored)
    demos: list[Demographics]
    reports: dict

    def demographics(self, source: str = "primary", age_offset: float = 0.0) -> pd.DataFrame:
        return demographics_table(self.demos, age_offset=age_offset, source=source)

    def save(self, cache_dir: Path) -> None:
        cache_dir.mkdir(parents=True, exist_ok=True)
        tag = f"{self.window_s}s"
        np.save(cache_dir / f"windows_{tag}.npy", self.windows)
        self.meta.to_csv(cache_dir / f"meta_{tag}.csv", index=False)
        (cache_dir / f"demos_{tag}.json").write_text(json.dumps([asdict(d) for d in self.demos], indent=1))
        (cache_dir / f"reports_{tag}.json").write_text(json.dumps(self.reports, indent=1, default=str))

    @classmethod
    def load(cls, cache_dir: Path, window_s: int) -> "WindowDataset":
        tag = f"{window_s}s"
        demos = [Demographics(**d) for d in json.loads((cache_dir / f"demos_{tag}.json").read_text())]
        return cls(
            window_s=window_s,
            n_samples=window_s * SAMPLE_RATE_HZ,
            meta=pd.read_csv(cache_dir / f"meta_{tag}.csv"),
            windows=np.load(cache_dir / f"windows_{tag}.npy"),
            demos=demos,
            reports=json.loads((cache_dir / f"reports_{tag}.json").read_text()),
        )


def window_bounds(
    row: LogRow,
    date_mmddyyyy: str,
    window_s: int,
    skip_windows: int,
) -> list[tuple[int, int, int]]:
    """``(j, start_ms, end_ms)`` for each kept window of a log row (in-lab.py:271-321).

    ``int((et - st).seconds / window_size)`` windows are laid back to back from the start time,
    the first ``skip_windows`` are dropped (upstream: ``60/window_size*2 - 1``, i.e. one window
    at 60 s, 7 at 15 s - the paper's "first two minutes" is not what the code does).
    """
    base = dt.datetime(1900, 1, 1)
    secs = int(((dt.datetime.combine(base, row.stop) - dt.datetime.combine(base, row.start)).seconds) / window_s)
    day = pd.Timestamp(dt.datetime.strptime(date_mmddyyyy, "%m/%d/%Y")).date()
    start = pd.Timestamp.combine(day, row.start)
    out = []
    for j in range(secs):
        if j >= skip_windows:
            s = start + pd.Timedelta(seconds=window_s * j)
            e = s + pd.Timedelta(seconds=window_s)
            out.append((j, local_to_epoch_ms(s), local_to_epoch_ms(e)))
    return out


def _slice(stream: WristStream, s_ms: int, e_ms: int) -> np.ndarray:
    i0 = int(np.searchsorted(stream.t_ms, s_ms, side="left"))
    i1 = int(np.searchsorted(stream.t_ms, e_ms, side="left"))
    return stream.x[i0:i1]


def _max_original_gap(orig_t: np.ndarray, s_ms: int, e_ms: int) -> float:
    """Largest spacing between original samples that touch the window (includes the neighbours)."""
    i0 = max(int(np.searchsorted(orig_t, s_ms, side="left")) - 1, 0)
    i1 = min(int(np.searchsorted(orig_t, e_ms, side="left")) + 1, orig_t.shape[0])
    seg = orig_t[i0:i1]
    return float(np.max(np.diff(seg))) if seg.shape[0] >= 2 else float("inf")


def build_window_dataset(
    data_root: Path,
    window_s: int = 60,
    *,
    participants: list[str] | None = None,
    skip_windows: int | None = None,
    verify_shipped: bool = True,
) -> WindowDataset:
    """All windows of all public participants with MetCart labels (in-lab.py:230-368)."""
    pdirs = sorted(p for p in data_root.glob("P*") if p.is_dir())
    if participants is not None:
        pdirs = [p for p in pdirs if p.name[1:] in participants]
    n_samples = window_s * SAMPLE_RATE_HZ
    skip = int(60 / window_s * 2 - 1) if skip_windows is None else skip_windows
    rows, win_list = [], []
    demos: list[Demographics] = []
    reports: dict = {"participants": {}, "skip_windows": skip, "window_s": window_s}
    for pdir in pdirs:
        pid = pdir.name[1:]
        start = MET_CART_START[pid]
        date = start[:10]
        demo = read_demographics(pdir)
        demos.append(demo)
        log = read_activity_log(pdir / "Calorie Harmony Phase 2 Activity Log.xlsx")
        met = read_metcart(pdir / "Metcart" / f"P{pid} (raw)", date)
        acc = read_wrist_stream(pdir, "Accelerometer", start)
        gyr = read_wrist_stream(pdir, "Gyroscope", start)
        rep = {
            "metcart_lines": [met.n_lines_total, met.n_data_like_lines, met.n_lines_parsed, int(met.met.shape[0])],
            "left_handed": demo.left_handed,
        }
        if verify_shipped:
            rep["shipped_acc"] = verify_against_shipped(pdir, "Accelerometer", acc)
            rep["shipped_gyr"] = verify_against_shipped(pdir, "Gyroscope", gyr)
        kept = skipped_not_included = problems = 0
        for row in log:
            if row.include != 1:
                skipped_not_included += 1
                continue
            if row.start is None or row.stop is None:
                problems += 1  # upstream: exception -> "encounters a problem" -> segment dropped
                continue
            for j, s_ms, e_ms in window_bounds(row, date, window_s, skip):
                g = _slice(gyr, s_ms, e_ms)
                a = _slice(acc, s_ms, e_ms)
                m_mask = (met.t_ms >= s_ms) & (met.t_ms < e_ms)
                m_vals = met.met[m_mask]
                # helper_preprocess.get_met_matcart: round(np.mean(...), 3) -> numpy rounding semantics
                met_raw = float(np.round(np.mean(m_vals), 3)) if m_vals.size else float("nan")
                buf = np.full((n_samples, 6), np.nan)
                buf[: min(g.shape[0], n_samples), 0:3] = g[:n_samples]
                buf[: min(a.shape[0], n_samples), 3:6] = a[:n_samples]
                win_list.append(buf)
                rows.append(
                    {
                        "participant": int(pid),
                        "activity": row.activity,
                        "segment": row.row,
                        "win_in_seg": j,
                        "start_ms": s_ms,
                        "n_gyr": g.shape[0],
                        "n_acc": a.shape[0],
                        "n_met_rows": int(m_vals.size),
                        "met_raw": met_raw,
                        "label": 0 if row.activity in SEDENTARY_ACTIVITIES else 1,
                        "max_gap_acc_ms": _max_original_gap(acc.original_t_ms, s_ms, e_ms),
                        "max_gap_gyr_ms": _max_original_gap(gyr.original_t_ms, s_ms, e_ms),
                        "left_handed": demo.left_handed,
                    }
                )
                kept += 1
        rep.update(windows=kept, segments_not_included=skipped_not_included, segments_with_problem=problems)
        reports["participants"][pid] = rep
    meta = pd.DataFrame(rows)
    # in-lab.py:335-337: METs below 1.0 are replaced by 1.0, NaN stays NaN.
    meta["met"] = np.where(meta["met_raw"] < 1.0, 1.0, meta["met_raw"])
    meta["activity_group"] = np.where(
        meta["activity"].isin(SEDENTARY_ACTIVITIES), "sedentary",
        np.where(meta["activity"].isin(LIGHT_ACTIVITIES), "light", "vigorous"),
    )
    return WindowDataset(
        window_s=window_s,
        n_samples=n_samples,
        meta=meta,
        windows=np.stack(win_list),
        demos=demos,
        reports=reports,
    )


# --------------------------------------------------------------------------------------
# Features exactly as the upstream pipeline computes them
# --------------------------------------------------------------------------------------

def upstream_window_arrays(ds: WindowDataset, *, mirror: bool = True) -> tuple[np.ndarray, np.ndarray]:
    """``(padded, mirrored_flag)``: windows with the upstream zero padding of short channels.

    ``get_train_data`` returns whatever samples fall in the window for each of the six series;
    short series are right-padded with zeros up to ``window_s * 20`` (in-lab.py:340-344).
    Left-handed participants are mirrored to a right-wrist frame when ``mirror`` is true.
    """
    w = ds.windows.copy()
    ng = ds.meta["n_gyr"].to_numpy()
    na = ds.meta["n_acc"].to_numpy()
    for i in np.nonzero((ng < ds.n_samples) | (na < ds.n_samples))[0]:
        w[i, ng[i] :, 0:3] = 0.0
        w[i, na[i] :, 3:6] = 0.0
    lh = ds.meta["left_handed"].to_numpy(bool)
    if mirror and lh.any():
        w[lh] = mirror_to_right_wrist(w[lh])
    return w, lh


def regression_intensity(ds: WindowDataset) -> np.ndarray:
    """(Intensity (Freq), Intensity (RMSSD_l1)) per window from the *unpadded* accelerometer samples."""
    out = np.full((len(ds.meta), 2), np.nan)
    na = ds.meta["n_acc"].to_numpy()
    for i in range(len(ds.meta)):
        n = int(min(na[i], ds.n_samples))
        if n >= 3:
            acc = ds.windows[i, :n, 3:6]
            out[i, 0] = freq_intensity(acc)
            out[i, 1] = rmssd_l1(acc)
    return out


def dataset_fingerprint(ds: WindowDataset) -> str:
    h = hashlib.sha256()
    h.update(np.ascontiguousarray(ds.windows).tobytes())
    h.update(ds.meta.to_csv(index=False).encode())
    return h.hexdigest()
